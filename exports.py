import io
import os
import re
from datetime import datetime
import pandas as pd
import openpyxl
from docx import Document


def generate_excel_bytes(data, sheet_name="Журнал"):
    """Генерує базовий Excel (XLSX) файл у вигляді байтів."""
    output = io.BytesIO()

    if isinstance(data, pd.DataFrame):
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            data.to_excel(writer, sheet_name=sheet_name, index=False)
    elif isinstance(data, dict):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        topics = data.get('topics', [])
        students = data.get('students', [])
        grades = data.get('grades', {})

        headers = ["Студент"] + [t.get('title', '') for t in topics]
        ws.append(headers)

        for student in students:
            row = [student]
            for topic in topics:
                title = topic.get('title', '')
                val = grades.get((student, title), "Не оцінено")
                row.append(val)
            ws.append(row)

        wb.save(output)

    output.seek(0)
    return output.getvalue()


def generate_docx_bytes(data, course_name="Журнал"):
    """Генерує Word документ (DOCX) з журналом."""
    doc = Document()
    doc.add_heading(f"Журнал успішності: {course_name}", level=1)

    if isinstance(data, pd.DataFrame):
        headers = list(data.columns)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        for idx, col_name in enumerate(headers):
            hdr_cells[idx].text = str(col_name)

        for _, row in data.iterrows():
            row_cells = table.add_row().cells
            for idx, col_name in enumerate(headers):
                val = row[col_name]
                row_cells[idx].text = str(val) if pd.notna(val) else ""

    elif isinstance(data, dict):
        topics = data.get('topics', [])
        students = data.get('students', [])
        grades = data.get('grades', {})

        table = doc.add_table(rows=1, cols=len(topics) + 1)
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Студент"
        for idx, topic in enumerate(topics):
            hdr_cells[idx + 1].text = topic.get('title', '')

        for student in students:
            row_cells = table.add_row().cells
            row_cells[0].text = student
            for idx, topic in enumerate(topics):
                title = topic.get('title', '')
                val = grades.get((student, title), "Не оцінено")
                row_cells[idx + 1].text = str(val) if val is not None else ""

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def generate_excel_template_full_bytes(course_name, teacher_name, journal_struct, template_path="Шаблон.xlsx"):
    """
    Заповнює обидва аркуші ('Оцінювання' та 'Теми занять') у файлі шаблону
    та повертає байтовий потік.
    """
    possible_paths = [template_path, "Шаблон. Оцінювання.xlsx", "Шаблон.xlsx"]
    actual_path = None
    for p in possible_paths:
        if os.path.exists(p):
            actual_path = p
            break

    if not actual_path:
        raise FileNotFoundError(f"Файл шаблону '{template_path}' не знайдено.")

    wb = openpyxl.load_workbook(actual_path)

    topics = journal_struct.get('topics', [])
    students = journal_struct.get('students', [])
    grades = journal_struct.get('grades', {})

    # --- 1. ЗАПОВНЕННЯ АРКУША "Оцінювання" ---
    if "Оцінювання" in wb.sheetnames:
        ws_grades = wb["Оцінювання"]
    else:
        ws_grades = wb.worksheets[0]

    ws_grades['A2'] = course_name

    col_idx = 3
    topic_cols = {}
    for topic in topics:
        title = topic.get('title', '')
        match = re.search(r'\((\d{4}-\d{2}-\d{2})\)', title)
        if match:
            date_str = match.group(1)
            try:
                date_val = datetime.strptime(date_str, '%Y-%m-%d').date()
            except ValueError:
                date_val = date_str
        else:
            date_val = ""

        cell = ws_grades.cell(row=8, column=col_idx)
        cell.value = date_val
        if isinstance(date_val, datetime) or hasattr(date_val, 'strftime'):
            cell.number_format = 'yyyy-mm-dd'

        topic_cols[title] = col_idx
        col_idx += 1

    for row_offset, student_name in enumerate(students):
        row_num = 9 + row_offset
        ws_grades.cell(row=row_num, column=1).value = row_offset + 1
        ws_grades.cell(row=row_num, column=2).value = student_name

        for topic in topics:
            title = topic.get('title', '')
            c_idx = topic_cols.get(title)
            if c_idx:
                val = grades.get((student_name, title))
                if val is not None and str(val).strip() != "" and str(val) != "Не оцінено":
                    try:
                        val_num = float(val) if '.' in str(val) else int(val)
                        ws_grades.cell(row=row_num, column=c_idx).value = val_num
                    except (ValueError, TypeError):
                        ws_grades.cell(row=row_num, column=c_idx).value = val

    # --- 2. ЗАПОВНЕННЯ АРКУША "Теми занять" ---
    if "Теми занять" in wb.sheetnames:
        ws_topics = wb["Теми занять"]
        ws_topics['D1'] = teacher_name

        for idx, topic in enumerate(topics, start=1):
            row_num = 3 + idx
            title_full = topic.get('title', '')
            clean_title = re.sub(r'\s*\(\d{4}-\d{2}-\d{2}\)', '', title_full).strip()

            match = re.search(r'\((\d{4}-\d{2}-\d{2})\)', title_full)
            if match:
                date_str = match.group(1)
                try:
                    date_val = datetime.strptime(date_str, '%Y-%m-%d').date()
                except ValueError:
                    date_val = date_str
            else:
                date_val = ""

            ws_topics.cell(row=row_num, column=1).value = idx
            cell_date = ws_topics.cell(row=row_num, column=2)
            cell_date.value = date_val
            if isinstance(date_val, datetime) or hasattr(date_val, 'strftime'):
                cell_date.number_format = 'yyyy-mm-dd'

            ws_topics.cell(row=row_num, column=4).value = clean_title

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()